import os
import io
from concurrent.futures import ThreadPoolExecutor
from pptx import Presentation
from PIL import Image

def extract_ppt_content(file_path: str) -> dict:
    """Extracts text content, notes, and OCRs embedded images from a PowerPoint presentation."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    from services.ocr_service import ocr_image
    
    prs = Presentation(file_path)
    total_slides = len(prs.slides)
    
    extracted_text = []
    pending_ocr = []
    with ThreadPoolExecutor(max_workers=2) as ocr_pool:
        for idx, slide in enumerate(prs.slides, start=1):
            slide_parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        line = paragraph.text.strip()
                        if line:
                            slide_parts.append(line)

                # OCR embedded images concurrently, with a small bounded queue.
                if getattr(shape, 'shape_type', None) == 13:
                    try:
                        image_bytes = shape.image.blob
                        img = Image.open(io.BytesIO(image_bytes)).convert("RGB")
                        part_index = len(slide_parts)
                        slide_parts.append(None)
                        pending_ocr.append((slide_parts, part_index, ocr_pool.submit(ocr_image, img)))
                        if len(pending_ocr) == 2:
                            _collect_ocr_results(pending_ocr)
                            pending_ocr.clear()
                    except Exception:
                        pass

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_parts.append(f"[Notes: {notes}]")

            if slide_parts:
                extracted_text.append((idx, slide_parts))

        if pending_ocr:
            _collect_ocr_results(pending_ocr)

    formatted_slides = []
    for idx, slide_parts in extracted_text:
        slide_content = "\n".join(part for part in slide_parts if part)
        if slide_content:
            formatted_slides.append(f"--- [Slide {idx}] ---\n{slide_content}")

    return {
        "type": "text",
        "total_units": total_slides,
        "unit_name": "slides",
        "content": "\n\n".join(formatted_slides)
    }


def _collect_ocr_results(pending):
    for slide_parts, part_index, future in pending:
        try:
            result = future.result()
            if result and "[OCR Failed" not in result:
                slide_parts[part_index] = f"[Image OCR: {result}]"
        except Exception:
            pass
