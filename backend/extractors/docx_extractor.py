import os
import docx

def extract_docx_content(file_path: str) -> dict:
    """Extracts text content, headings, and tables from a Word (.docx) document."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    doc = docx.Document(file_path)
    paragraphs = []
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            if p.style and p.style.name and p.style.name.startswith("Heading"):
                paragraphs.append(f"## {text}")
            else:
                paragraphs.append(text)
                
    # Also extract tables
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(row_cells):
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            paragraphs.append("\n[Table Data]:\n" + "\n".join(table_rows))
            
    full_text = "\n\n".join(paragraphs)
    return {
        "type": "text",
        "total_units": max(len(doc.paragraphs), 1),
        "unit_name": "paragraphs",
        "content": full_text
    }
