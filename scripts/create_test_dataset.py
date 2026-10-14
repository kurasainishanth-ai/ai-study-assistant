"""Generate deterministic synthetic multi-format StudyLane test materials.

Requires the optional libraries already used by the project environment:
PyMuPDF (fitz), python-docx, python-pptx, and Pillow.
"""

from pathlib import Path

import pymupdf as fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "sample_data" / "test_dataset" / "formats"


def font(size=24, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def make_chart(path: Path):
    image = Image.new("RGB", (1100, 700), "white")
    draw = ImageDraw.Draw(image)
    title_font, label_font = font(38, True), font(25)
    draw.text((180, 28), "Algorithm Runtime Test (synthetic)", fill="#111111", font=title_font)
    x0, y0, x1, y1 = 120, 590, 1030, 130
    draw.line((x0, y0, x0, y1), fill="#111111", width=4)
    draw.line((x0, y0, x1, y0), fill="#111111", width=4)
    draw.text((25, 340), "Time (ms)", fill="#111111", font=label_font)
    draw.text((500, 625), "Input size n", fill="#111111", font=label_font)
    points = [(220, 530), (390, 455), (560, 380), (730, 305), (900, 230)]
    draw.line(points, fill="#1464b4", width=8)
    for x, y in points:
        draw.ellipse((x - 10, y - 10, x + 10, y + 10), fill="#1464b4")
    for x, label in zip([220, 390, 560, 730, 900], [100, 200, 300, 400, 500]):
        draw.text((x - 25, 600), str(label), fill="#111111", font=label_font)
    draw.text((740, 165), "Observed trend rises with n", fill="#111111", font=label_font)
    image.save(path)


def make_diagram(path: Path):
    image = Image.new("RGB", (1400, 520), "white")
    draw = ImageDraw.Draw(image)
    title_font, label_font = font(38, True), font(27, True)
    draw.text((420, 25), "Neural Network Forward Pass", fill="#111111", font=title_font)
    boxes = [(45, 205, 330, 340), (420, 205, 710, 340), (800, 205, 1080, 340), (1170, 205, 1360, 340)]
    labels = ["Input features", "Weighted sum\nz = w*x + b", "ReLU\nmax(0, z)", "Output"]
    colors = ["#dceeff", "#e9e0ff", "#e2f4e8", "#fff0d9"]
    for box, label, color in zip(boxes, labels, colors):
        draw.rounded_rectangle(box, radius=24, fill=color, outline="#333333", width=4)
        draw.multiline_text((box[0] + 24, box[1] + 35), label, fill="#111111", font=label_font, spacing=12, align="center")
    for left, right in zip(boxes, boxes[1:]):
        y = (left[1] + left[3]) // 2
        draw.line((left[2] + 8, y, right[0] - 12, y), fill="#333333", width=6)
        draw.polygon([(right[0] - 12, y), (right[0] - 34, y - 15), (right[0] - 34, y + 15)], fill="#333333")
    image.save(path)


def make_pdf(chart_path: Path):
    path = OUT / "algorithm_runtime.pdf"
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((54, 55), "Algorithm Runtime — Synthetic Lab Sheet", fontsize=20)
    page.insert_text((54, 100), "[Page 1]", fontsize=12)
    page.insert_textbox(fitz.Rect(54, 125, 550, 220),
                        "This synthetic lab compares input size with measured runtime. The chart is illustrative, not a formal complexity proof. Binary search requires sorted input and performs logarithmically many comparisons in the worst case.", fontsize=12)
    page.insert_image(fitz.Rect(68, 250, 545, 552), filename=str(chart_path))
    page.insert_text((54, 590), "Figure 1: Observed runtime trend across five input sizes.", fontsize=11)
    page2 = doc.new_page(width=612, height=792)
    page2.insert_text((54, 55), "[Page 2]", fontsize=12)
    page2.insert_textbox(fitz.Rect(54, 90, 550, 250),
                         "Interpretation task: identify the horizontal axis, vertical axis, and overall trend. Do not infer an exact asymptotic bound from this chart alone. A controlled benchmark needs repeated trials and documented hardware.", fontsize=12)
    doc.save(path)
    doc.close()
    return path


def make_docx():
    path = OUT / "study_methods.docx"
    doc = Document()
    doc.add_heading("Study Methods — Synthetic Handout", 0)
    doc.add_heading("[Section 1] Spaced practice", level=1)
    doc.add_paragraph("Spaced practice revisits material in multiple sessions with time between sessions. A schedule should fit the learner's available time and should not be changed without their agreement.")
    doc.add_heading("[Section 2] Retrieval practice", level=1)
    doc.add_paragraph("Retrieval practice asks the learner to recall an idea before looking at the answer. Feedback helps identify what to review. One quiz result is limited evidence of mastery.")
    doc.add_heading("[Section 3] Synthetic checklist table", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Practice action"
    table.rows[0].cells[1].text = "Purpose"
    for action, purpose in [("Recall before checking", "Practice retrieval"), ("Return after a delay", "Space practice"), ("Review the mistaken step", "Use feedback")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = action, purpose
    doc.save(path)
    return path


def make_pptx(diagram_path: Path):
    path = OUT / "neural_network_slides.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Neural Network Forward Pass"
    textbox = slide.shapes.add_textbox(Inches(0.7), Inches(1.25), Inches(8.8), Inches(0.7))
    textbox.text_frame.text = "[Slide 1] A neuron combines input values using learned weights and a bias."
    slide.shapes.add_picture(str(diagram_path), Inches(0.55), Inches(2.0), width=Inches(9.2))
    slide2 = deck.slides.add_slide(deck.slide_layouts[5])
    slide2.shapes.title.text = "Activation and example"
    box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.5), Inches(2.5))
    frame = box.text_frame
    frame.text = "[Slide 2] ReLU returns max(0, z)."
    p = frame.add_paragraph()
    p.text = "For z = -2, the output is 0. For z = 5, the output is 5."
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(35, 55, 80)
    p = frame.add_paragraph()
    p.text = "Check: identify the operation that occurs before activation."
    p.font.size = Pt(20)
    deck.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    chart_path = OUT / "runtime_chart.png"
    diagram_path = OUT / "neural_network_diagram.png"
    make_chart(chart_path)
    make_diagram(diagram_path)
    outputs = [make_pdf(chart_path), make_docx(), make_pptx(diagram_path), chart_path, diagram_path]
    print("Created synthetic format test assets:")
    for item in outputs:
        print(f"- {item.relative_to(ROOT)} ({item.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
