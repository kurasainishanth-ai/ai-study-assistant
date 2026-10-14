import json
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from extractors.dispatcher import is_supported_file  # noqa: E402


class TestSyntheticDatasetAssets(unittest.TestCase):
    def test_manifest_files_exist_and_formats_are_supported(self):
        manifest_path = ROOT / "sample_data" / "test_dataset" / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.assertEqual(len(manifest["materials"]), 8)
        for material in manifest["materials"]:
            path = manifest_path.parent / material["file"]
            self.assertTrue(path.is_file(), str(path))
            self.assertTrue(is_supported_file(str(path)), str(path))

    def test_generated_binary_assets_parse_without_calling_ocr_service(self):
        import pymupdf as fitz
        from docx import Document
        from pptx import Presentation
        from PIL import Image

        formats = ROOT / "sample_data" / "test_dataset" / "formats"
        pdf = fitz.open(formats / "algorithm_runtime.pdf")
        self.assertEqual(len(pdf), 2)
        self.assertIn("binary search", pdf[0].get_text().lower())
        pdf.close()

        word = Document(formats / "study_methods.docx")
        self.assertTrue(any("retrieval practice" in p.text.lower() for p in word.paragraphs))
        self.assertEqual(len(word.tables), 1)

        slides = Presentation(formats / "neural_network_slides.pptx")
        self.assertEqual(len(slides.slides), 2)
        visible_text = " ".join(shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("ReLU", visible_text)
        self.assertTrue(any(shape.shape_type == 13 for shape in slides.slides[0].shapes))

        for filename in ("runtime_chart.png", "neural_network_diagram.png"):
            with Image.open(formats / filename) as image:
                self.assertGreater(image.width, 500)
                self.assertGreater(image.height, 200)


if __name__ == "__main__":
    unittest.main()
