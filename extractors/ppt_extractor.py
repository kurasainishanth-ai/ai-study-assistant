import os
import io
from pptx import Presentation
from PIL import Image

def extract_ppt_content(file_path: str) -> dict:
    """Extracts text content, notes, and OCRs embedded images from a PowerPoint presentation."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    from services.ocr_service import ocr_image
    
    prs = Presentation(file_path)
    total_slides = len(prs.slides)
    
    extracted_text = []
    for idx, slide in enumerate(prs.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_parts.append(line)
            
            # Check for images (msoPICTURE == 13)
            if getattr(shape, 'shape_type', None) == 13:
                try:
                    image_bytes = shape.image.blob
                    img = Image.open(io.BytesIO(image_bytes))
                    ocr_result = ocr_image(img)
                    if ocr_result and "[OCR Failed" not in ocr_result:
                        slide_parts.append(f"[Image OCR: {ocr_result}]")
                except Exception as e:
                    pass
        
        # Check for slide notes if available
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_parts.append(f"[Notes: {notes}]")
                
        if slide_parts:
            slide_content = "\n".join(slide_parts)
            extracted_text.append(f"--- [Slide {idx}] ---\n{slide_content}")
            
    full_text = "\n\n".join(extracted_text)
    return {
        "type": "text",
        "total_units": total_slides,
        "unit_name": "slides",
        "content": full_text
    }
